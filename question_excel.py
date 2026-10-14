from io import BytesIO
from pathlib import Path
import sqlite3

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook


router = APIRouter(prefix="/api/questions", tags=["Question Excel"])
DB_PATH = Path(__file__).resolve().parent / "career_course.db"

HEADERS = [
    "question_text", "category",
    "option_a", "option_b", "option_c", "option_d",
    "score_a_category", "score_a_points",
    "score_b_category", "score_b_points",
    "score_c_category", "score_c_points",
    "score_d_category", "score_d_points",
    "is_active",
]


def make_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(HEADERS)

    for row in rows:
        sheet.append(row)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.get("/template")
def download_question_template():
    example = [
        "I enjoy solving computer problems.",
        "Technology",
        "Strongly agree", "Agree", "Disagree", "Strongly disagree",
        "Technology", 4,
        "Technology", 3,
        "Technology", 2,
        "Technology", 1,
        1,
    ]

    output = make_workbook([example])

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="careerpath_questions_template.xlsx"'},
    )


@router.get("/export")
def export_questions():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row

    try:
        questions = connection.execute(
            "SELECT * FROM questions ORDER BY id"
        ).fetchall()

        rows = []

        for question in questions:
            question_id = question["id"]

            scores = connection.execute(
                """
                SELECT option_code, interest_category, points
                FROM question_option_scores
                WHERE question_id = ?
                """,
                (question_id,),
            ).fetchall()

            score_map = {
                score["option_code"]: (
                    score["interest_category"],
                    score["points"],
                )
                for score in scores
            }

            row = [
                question["question_text"],
                question["category"],
                question["option_a"],
                question["option_b"],
                question["option_c"],
                question["option_d"],
            ]

            for code in ["A", "B", "C", "D"]:
                category, points = score_map.get(code, ("", 0))
                row.extend([category, points])

            row.append(question["is_active"])
            rows.append(row)

        output = make_workbook(rows)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="careerpath_questions.xlsx"'},
        )

    finally:
        connection.close()


@router.post("/import")
async def import_questions(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Upload an Excel .xlsx file."
        )

    try:
        content = await file.read()
        workbook = load_workbook(BytesIO(content), data_only=True)
        sheet = workbook.active
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to read this Excel file. Please use the CareerPath template."
        )

    header = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in sheet[1]
    ]

    missing = [name for name in HEADERS if name not in header]

    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required Excel columns: {', '.join(missing)}"
        )

    positions = {name: header.index(name) for name in HEADERS}
    connection = sqlite3.connect(DB_PATH)
    imported = 0
    errors = []

    try:
        for row_number, excel_row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not any(value is not None for value in excel_row):
                continue

            def value(name):
                index = positions[name]
                return excel_row[index] if index < len(excel_row) else None

            question_text = str(value("question_text") or "").strip()
            category = str(value("category") or "").strip()
            options = [
                str(value("option_a") or "").strip(),
                str(value("option_b") or "").strip(),
                str(value("option_c") or "").strip(),
                str(value("option_d") or "").strip(),
            ]

            if not question_text or not category or not all(options):
                errors.append(
                    f"Row {row_number}: question, category, and all four options are required."
                )
                continue

            try:
                connection.execute("SAVEPOINT import_row")

                active_value = value("is_active")
                is_active = 1 if active_value in (None, 1, True, "1", "Yes", "yes") else 0

                cursor = connection.execute(
                    """
                    INSERT INTO questions
                    (question_text, category, option_a, option_b,
                     option_c, option_d, is_active)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        question_text, category,
                        options[0], options[1],
                        options[2], options[3],
                        is_active,
                    ),
                )

                question_id = cursor.lastrowid

                for code in ["A", "B", "C", "D"]:
                    category_value = str(
                        value(f"score_{code.lower()}_category") or ""
                    ).strip()
                    points_value = value(f"score_{code.lower()}_points")

                    if not category_value:
                        continue

                    points = int(points_value or 0)

                    connection.execute(
                        """
                        INSERT INTO question_option_scores
                        (question_id, option_code, interest_category, points)
                        VALUES (?, ?, ?, ?)
                        """,
                        (question_id, code, category_value, points),
                    )

                connection.execute("RELEASE SAVEPOINT import_row")
                imported += 1

            except Exception as error:
                connection.execute("ROLLBACK TO SAVEPOINT import_row")
                connection.execute("RELEASE SAVEPOINT import_row")
                errors.append(f"Row {row_number}: {error}")

        connection.commit()

        return {
            "message": "Question import finished. Existing questions were preserved.",
            "imported": imported,
            "errors": errors,
        }

    except Exception as error:
        connection.rollback()
        raise HTTPException(status_code=500, detail=str(error))

    finally:
        connection.close()